# -*- coding: utf-8 -*-
"""标准答案校验（同格式输出与分析视图共用同一口径）。

支持两种输出：
  - 同格式：表头第 1 行，A 列序号 + M/N；用 Sheet 名作为机型文件。
  - 分析视图：表头含 机型文件/原序号 + M/N；读取机型代码。

判定：
  ok    = 有对应标准答案，且 M/N 与标准答案一致
  bad   = 有对应标准答案，M/N 非空但不一致
  blank = 有对应标准答案，但 M/N 为空（未分类/未完成）
  missing = 标准答案中没有该序号（不计入一致/不符）
"""
import sys

from openpyxl import load_workbook

from sip_rules import header_key


def parse_seq(v):
    if isinstance(v, (int, float)):
        return int(v)
    if isinstance(v, str) and v.strip().isdigit():
        return int(v.strip())
    return None


def read_output_records(path):
    """返回 [{sheet,file,model,seq,main,sub}]。"""
    wb = load_workbook(path, data_only=True)
    records = []
    for ws in wb.worksheets:
        if ws.title == '说明' or ws.title.startswith('待确认'):
            continue
        hr = None
        for r in range(1, min(ws.max_row, 15) + 1):
            vals = [header_key(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
            if '主顺序' in vals and '子顺序' in vals and (
                    '序号' in vals or '原序号' in vals):
                hr = r
                break
        if hr is None:
            continue
        col = {}
        for c in range(1, ws.max_column + 1):
            v = header_key(ws.cell(hr, c).value)
            if v and v not in col:
                col[v] = c
        c_main = col['主顺序']
        c_sub = col['子顺序']

        if '机型文件' in col and '原序号' in col:
            c_file = col['机型文件']
            c_seq = col['原序号']
            c_model = col.get('机型代码')
            for r in range(hr + 1, ws.max_row + 1):
                f = ws.cell(r, c_file).value
                seq = parse_seq(ws.cell(r, c_seq).value)
                if not f or seq is None:
                    continue
                records.append({
                    'sheet': ws.title,
                    'file': str(f),
                    'model': (str(ws.cell(r, c_model).value).strip()
                              if c_model and ws.cell(r, c_model).value else ''),
                    'seq': seq,
                    'main': str(ws.cell(r, c_main).value or '').strip(),
                    'sub': str(ws.cell(r, c_sub).value or '').strip(),
                })
        elif '序号' in col:
            c_seq = col['序号']
            for r in range(hr + 1, ws.max_row + 1):
                seq = parse_seq(ws.cell(r, c_seq).value)
                if seq is None:
                    continue
                records.append({
                    'sheet': ws.title,
                    'file': ws.title,
                    'model': '',
                    'seq': seq,
                    'main': str(ws.cell(r, c_main).value or '').strip(),
                    'sub': str(ws.cell(r, c_sub).value or '').strip(),
                })
    return records


def verify_records(records, ref):
    """按标准答案校验，返回每个机型的结果列表。"""
    grouped = {}
    for rec in records:
        grouped.setdefault(rec['file'], []).append(rec)

    targets = []
    for f, recs in grouped.items():
        models = {r['model'] for r in recs if r['model']}
        if ref.get('sheet_title') and ref['sheet_title'].lower() in f.lower():
            targets.append(f)
        elif ref.get('model') and ref['model'] in models:
            targets.append(f)

    results = []
    for f in sorted(targets):
        ok = bad = blank = missing = 0
        diffs = []
        for rec in grouped[f]:
            exp = ref['seq'].get(rec['seq'])
            if exp is None:
                missing += 1
                continue
            got = (rec['main'], rec['sub'])
            if not got[0] or not got[1]:
                blank += 1
                continue
            if got == exp:
                ok += 1
            else:
                bad += 1
                if len(diffs) < 5:
                    diffs.append((rec['seq'], got, exp))
        results.append({'file': f, 'ok': ok, 'bad': bad,
                        'blank': blank, 'missing': missing, 'diffs': diffs})
    return results


def format_results(results):
    lines = []
    if not results:
        return ['标准答案比对: 没有匹配的机型文件（未完成校验）']
    for r in results:
        if r['ok'] + r['bad'] == 0:
            if r['blank']:
                lines.append(f"标准答案比对 {r['file']}: 无可比对行；"
                             f"{r['blank']} 行 M/N 为空（未完成校验）")
            else:
                lines.append(f"标准答案比对 {r['file']}: 没有可比序号（未完成校验）")
            continue
        line = (f"标准答案比对 {r['file']}: {r['ok']}/{r['ok'] + r['bad']} 一致, "
                f"{r['bad']} 不符")
        if r['blank']:
            line += f", {r['blank']} 行 M/N 为空（未完成）"
        lines.append(line)
        for d in r['diffs']:
            lines.append(f'   {d}')
    return lines


def verify_file(output_path, ref):
    return verify_records(read_output_records(output_path), ref)


def exit_code(results):
    if not results:
        return 2
    total_ok = sum(r['ok'] for r in results)
    total_bad = sum(r['bad'] for r in results)
    total_blank = sum(r['blank'] for r in results)
    if total_ok + total_bad == 0:
        return 2
    return 1 if (total_bad or total_blank) else 0


if __name__ == '__main__':
    if len(sys.argv) != 3:
        print('用法: python -I sip_verify.py <输出.xlsx> <标准答案.xlsx>', file=sys.stderr)
        sys.exit(2)
    from sip_rules import load_reference
    try:
        sys.stdout.reconfigure(encoding='utf-8', errors='replace')
    except Exception:
        pass
    results = verify_file(sys.argv[1], load_reference(sys.argv[2]))
    print('\n'.join(format_results(results)))
    sys.exit(exit_code(results))
