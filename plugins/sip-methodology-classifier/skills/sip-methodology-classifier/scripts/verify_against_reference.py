# -*- coding: utf-8 -*-
"""用标准答案校验已生成的对应表。

用法：
  python -I verify_against_reference.py --output <对应表.xlsx> --reference <标准答案.xlsx>

匹配规则（与 build_sip_mapping.py 一致）：
  - 标准答案的工作表名出现在输出“机型文件”名中，或
  - 输出行的“机型代码”与标准答案的机型代码一致。

输出：逐行比对主顺序/子顺序，打印一致/不一致数量与前若干条差异；
      匹配到文件但没有任何可比序号时，明确显示“未完成校验”。

退出码：全部一致 0；有不一致 1；无法比对 / 未完成校验 2。
"""
import argparse
import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from openpyxl import load_workbook
from sip_rules import load_reference, header_key

try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass


def read_output(path):
    """读取输出工作簿，兼容两种格式。

    - 同格式：表头第 1 行，A 列序号 + M/N；用 Sheet 名作为机型文件。
    - 分析视图：表头含 机型文件/原序号 + M/N；按行读取机型代码。

    返回 {机型文件: {'model': 机型代码, 'rows': {原序号: (主顺序, 子顺序)}}}。
    """
    wb = load_workbook(path, data_only=True)
    out = {}
    for ws in wb.worksheets:
        if ws.title == '说明' or ws.title.startswith('待确认'):
            continue
        hr = None
        for r in range(1, min(ws.max_row, 15) + 1):
            vals = [header_key(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
            if '主顺序' in vals and '子顺序' in vals and (
                    '序号' in vals or '原序号' in vals):
                hr = r
                break
        if hr is None:
            continue
        col = {}
        for c in range(1, ws.max_column + 1):
            v = header_key(ws.cell(hr, c).value)
            if v and v not in col:
                col[v] = c
        c_main = col['主顺序']
        c_sub = col['子顺序']

        if '机型文件' in col and '原序号' in col:
            # 分析视图
            c_file = col['机型文件']
            c_seq = col['原序号']
            c_model = col.get('机型代码')
            for r in range(hr + 1, ws.max_row + 1):
                f = ws.cell(r, c_file).value
                seq = ws.cell(r, c_seq).value
                main = ws.cell(r, c_main).value
                sub = ws.cell(r, c_sub).value
                if not f or not isinstance(seq, (int, float)) or not main or not sub:
                    continue
                rec = out.setdefault(str(f), {'model': '', 'rows': {}})
                if c_model:
                    mv = ws.cell(r, c_model).value
                    if mv is not None and str(mv).strip():
                        rec['model'] = str(mv).strip()
                rec['rows'][int(seq)] = (str(main).strip(), str(sub).strip())
        elif '序号' in col:
            # 同格式输出
            c_seq = col['序号']
            rec = out.setdefault(ws.title, {'model': '', 'rows': {}})
            for r in range(hr + 1, ws.max_row + 1):
                seq = ws.cell(r, c_seq).value
                main = ws.cell(r, c_main).value
                sub = ws.cell(r, c_sub).value
                if not isinstance(seq, (int, float)) or not main or not sub:
                    continue
                rec['rows'][int(seq)] = (str(main).strip(), str(sub).strip())
    return out


def main():
    ap = argparse.ArgumentParser(description='用标准答案校验已生成的对应表')
    ap.add_argument('--output', required=True, help='生成的对应表 xlsx')
    ap.add_argument('--reference', required=True,
                    help='标准答案 xlsx（第一个工作表，表头含 序号/检验工序/检验项目/主顺序/子顺序）')
    args = ap.parse_args()

    ref = load_reference(args.reference)
    out = read_output(args.output)

    targets = []
    for f, rec in out.items():
        if ref.get('sheet_title') and ref['sheet_title'].lower() in f.lower():
            targets.append(f)
        elif ref.get('model') and rec.get('model') == ref['model']:
            targets.append(f)
    if not targets:
        print('无法比对：输出中没有与标准答案的工作表名或机型代码匹配的机型文件')
        return 2

    total_ok = total_bad = total_compared = 0
    for f in sorted(targets):
        ok = bad = 0
        diffs = []
        for seq, got in sorted(out[f]['rows'].items()):
            exp = ref['seq'].get(seq)
            if exp is None:
                continue
            if got == exp:
                ok += 1
            else:
                bad += 1
                if len(diffs) < 10:
                    diffs.append((seq, got, exp))
        if ok + bad == 0:
            print(f'{f}: 匹配到文件，但没有可比的序号行（未完成校验）')
            continue
        total_compared += ok + bad
        total_ok += ok
        total_bad += bad
        print(f'{f}: {ok}/{ok + bad} 一致, {bad} 不符')
        for d in diffs:
            print('  ', d)

    if total_compared == 0:
        print('未完成校验：匹配到文件，但没有任何可比的序号行')
        return 2
    print(f'合计: {total_ok}/{total_compared} 一致, {total_bad} 不符')
    return 0 if total_bad == 0 else 1


if __name__ == '__main__':
    sys.exit(main())
