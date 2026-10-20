# -*- coding: utf-8 -*-
"""按 SIP 方法论两层顺序模型归类《成品检验指导书》。

用法：
  python -I build_sip_mapping.py --input <zip或目录> --output <输出.xlsx> \
      [--reference <标准答案.xlsx>] [--project <单项目名称>] [--keep-temp]

输入：
  - zip（文件名可能是 GBK）或文件夹；
  - 多项目：一级子文件夹 = 项目；单项目：xlsx 放在根目录并用 --project 命名。
输出：
  - 说明、每个项目一个 Sheet、待确认项（规则推断）。
"""
import argparse
import os
import shutil
import sys
import tempfile

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

import sip_rules
from sip_rules import (MAIN_ORDER, SUB_ORDER, MAIN_IDX, SUB_IDX,
                       norm, header_key, classify, load_reference)
from extract_gbk_zip import extract as extract_zip

try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

EXCLUDE_HEADERS = {'检验频率', '重要等级', '检验记录表', '备注'}
SUB_MEANING = {
    sip_rules.M: [('1级', '受订单核对、机台名称/编号与铭牌核对、生产计划核对'),
                  ('2级', '加密狗规格、工控机序列号、配件/清单核对')],
    sip_rules.W: [('1级', '整机外观、履历表、水槽盖板、压力表、地脚、接头、颜色、顶丝、光纤、离子棒安装、外罩、清洁、出货防护'),
                  ('2级', '可动件/内部布局：手动操作、切割门、拖链线、网线、运动部件线材、强弱电分离'),
                  ('3级', '部件外观：主轴冷却水接口、主轴轴心、屏幕后盖板、防尘罩、NCS固定钣金、热扩台盘调节')],
    sip_rules.D: [('1级', '厂务动力与基础：供电电压、水压、水平、压缩空气、排风、I/O点位'),
                  ('2级', 'Y轴重复/定位精度、精度补偿'), ('3级', '补偿文件'),
                  ('4级', '精度检查表：直线度、垂直度、跳动、平整度、重复精度、过冲'),
                  ('5级', '功能单元：测高、BBD、规整、搬运/复位、清洗、UV、空切、吸盘、单元动作、加热/冷扩/热扩、首料检'),
                  ('6级', '料库检料传感器'), ('7级', '视觉对位、相机验证/对位、X/Y轴垂直度'),
                  ('8级', '气缸示教'), ('9级', '静电、接地、离子棒/离子风扇、漏电、电子断路器'),
                  ('10级', '主轴转速、设备噪音'), ('11级', '联锁、设备序列号、限位片'),
                  ('12级', '安全防护、版本、条码枪/扫码器、法兰刀片、整机')],
    sip_rules.Q: [('1级', '气源与空气过滤组合'), ('2级', '吹气流量、气帘、气枪、气幕'),
                  ('3级', '真空、真空发生器、气密'), ('4级', '气管')],
    sip_rules.S: [('1级', '水流量、二流体、水温传感器、BBD/NCS传感器、水帘/气帘角度、水气帘高度、隔板、清洗过程'),
                  ('2级', '喷水架位置'), ('3级', '漏水检测、漏水绳、发泡机、防水防锈、切割门钣金漏水、水雾')],
    sip_rules.L: [('1级', '老化前确认：传送精度、气压阈值、电机速度力矩、位置居中'),
                  ('2级', '老化运行：IO/复位/急停/报警灯/各轴复位/8小时老化/功能流程')],
    sip_rules.G: [('1级', '切割精度验证'), ('2级', '切割品质验证'),
                  ('3级', '刀痕宽度、UV膜刀痕、Dummy切深')],
    sip_rules.U: [('1级', '断电验证'), ('2级', '断电影响确认（工控机）')],
    sip_rules.B: [('1级', '加密狗永久版本/序列号'), ('2级', 'Windows更新/睡眠/激活设置'),
                  ('3级', '机台备份'), ('4级', '备份后软件确认'), ('5级', '杀毒软件安装'),
                  ('6级', '查杀病毒记录'), ('7级', '加密执行确认'), ('8级', 'USB接口确认')],
}


# ---------------------------------------------------------------- 单元格工具
def is_yellow(cell):
    f = cell.fill
    if f is None or f.fill_type is None:
        return False
    c = f.start_color
    rgb = getattr(c, 'rgb', None)
    if isinstance(rgb, str) and len(rgb) == 8:
        r, g, b = int(rgb[2:4], 16), int(rgb[4:6], 16), int(rgb[6:8], 16)
        return r >= 200 and g >= 180 and b <= 160
    return False


def cell_text(ws, row, col):
    if col is None:
        return ''
    cell = ws.cell(row, col)
    if is_yellow(cell):
        return ''
    v = cell.value
    return str(v).strip() if v is not None else ''


def find_header_row(ws):
    for r in range(1, min(ws.max_row, 15) + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if '序号' in vals and '检验工序' in vals and '检验项目' in vals:
            return r
    raise ValueError(f'{ws.title}: 找不到表头行（需要 序号/检验工序/检验项目）')


def get_columns(ws, header_row):
    cmap = {}
    for c in range(1, ws.max_column + 1):
        v = header_key(ws.cell(header_row, c).value)
        if v and v not in cmap:
            cmap[v] = c
    missing = [x for x in ('序号', '检验工序', '检验项目') if x not in cmap]
    if missing:
        raise ValueError(f'{ws.title}: 缺少必需列 {missing}')
    return cmap


# ---------------------------------------------------------------- 解析
def parse_sheet(path, project):
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    hr = find_header_row(ws)
    cmap = get_columns(ws, hr)
    c_seq = cmap['序号']
    c_stage = cmap['检验工序']
    c_item = cmap['检验项目']
    c_std = cmap.get('检验标准')
    c_method = cmap.get('检验方法')
    c_tool = cmap.get('检验工具')

    sub_cols = []
    if c_std:
        sub_cols = [c for c in range(c_item + 1, c_std) if c != c_stage]

    product = model = ''
    for r in range(1, hr):
        for c in range(1, ws.max_column + 1):
            v = cell_text(ws, r, c)
            if v.startswith('产品名称'):
                product = v.split('：', 1)[-1].strip()
            elif v.startswith('机型代码'):
                model = v.split('：', 1)[-1].strip()

    rows = []
    stage = item = ''
    skipped_yellow = 0
    for r in range(hr + 1, ws.max_row + 1):
        seqv = ws.cell(r, c_seq).value
        if isinstance(seqv, (int, float)):
            seq = int(seqv)
        elif isinstance(seqv, str) and seqv.strip().isdigit():
            seq = int(seqv.strip())
        else:
            continue
        if is_yellow(ws.cell(r, c_seq)):
            skipped_yellow += 1
            continue
        b = cell_text(ws, r, c_stage)
        if b:
            stage = norm(b)
            item = ''
        c = cell_text(ws, r, c_item)
        if c:
            item = norm(c)
        sub = ' / '.join(x for x in (cell_text(ws, r, sc) for sc in sub_cols) if x)
        rows.append(dict(
            project=project,
            file=os.path.splitext(os.path.basename(path))[0],
            product=product,
            model=model,
            seq=seq,
            stage=stage,
            item=item,
            sub=sub,
            std=cell_text(ws, r, c_std),
            method=cell_text(ws, r, c_method),
            tool=cell_text(ws, r, c_tool),
            src=r,
        ))
    return rows, skipped_yellow


# ---------------------------------------------------------------- 项目发现
def _xlsx_files(d):
    return [f for f in sorted(os.listdir(d))
            if f.lower().endswith('.xlsx') and not f.startswith('~$')]


def _data_root(d):
    if _xlsx_files(d):
        return d
    subs = [x for x in sorted(os.listdir(d)) if os.path.isdir(os.path.join(d, x))]
    if len(subs) == 1:
        return _data_root(os.path.join(d, subs[0]))
    return d


def discover_projects(root, project_name=None):
    root = _data_root(root)
    projects = []
    for entry in sorted(os.listdir(root)):
        p = os.path.join(root, entry)
        if os.path.isdir(p):
            files = _xlsx_files(p)
            if files:
                projects.append((entry, p, files))
    root_files = _xlsx_files(root)
    if root_files:
        name = project_name or os.path.basename(os.path.abspath(root)) or '默认项目'
        projects.append((name, root, root_files))
    if not projects:
        raise ValueError(f'在 {root} 下没有找到 .xlsx 文件')
    return projects


# ---------------------------------------------------------------- 校验
def verify_reference(all_rows, ref):
    if not ref:
        return []
    targets = set()
    for row in all_rows:
        stem = row['file'].lower()
        if ref.get('sheet_title') and ref['sheet_title'].lower() in stem:
            targets.add(row['file'])
        elif ref.get('model') and row['model'] == ref['model']:
            targets.add(row['file'])
    out = []
    for f in sorted(targets):
        rows = [r for r in all_rows if r['file'] == f]
        ok = bad = 0
        diffs = []
        for r in rows:
            exp = ref['seq'].get(r['seq'])
            if exp is None:
                continue
            if (r['main'], r['sublevel']) == exp:
                ok += 1
            else:
                bad += 1
                if len(diffs) < 8:
                    diffs.append((r['seq'], (r['main'], r['sublevel']), exp))
        out.append((f, ok, bad, diffs))
    return out


# ---------------------------------------------------------------- 输出
def safe_sheet_name(name, used):
    for ch in '\\/*?:[]':
        name = name.replace(ch, '_')
    name = (name or '项目')[:31]
    base = name
    i = 2
    while name in used:
        suffix = f'_{i}'
        name = (base[:31 - len(suffix)] + suffix)
        i += 1
    used.add(name)
    return name


def write_workbook(all_rows, projects, out_path, ref=None, verify=None,
                   skipped_yellow=0):
    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style='thin', color='BFBFBF')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    head_fill = PatternFill('solid', fgColor='305496')
    head_font = Font(name='微软雅黑', size=10, bold=True, color='FFFFFF')
    body_font = Font(name='微软雅黑', size=10)
    main_fill = PatternFill('solid', fgColor='1F3864')
    main_font = Font(name='微软雅黑', size=12, bold=True, color='FFFFFF')
    sub_fill = PatternFill('solid', fgColor='D9E1F2')
    sub_font = Font(name='微软雅黑', size=10, bold=True, color='1F3864')
    wrap = Alignment(vertical='center', wrap_text=True)
    center = Alignment(horizontal='center', vertical='center', wrap_text=True)

    headers = ['主顺序', '子顺序', '机型文件', '产品名称', '机型代码', '原序号',
               '检验工序', '检验项目', '子项', '检验标准', '检验方法', '检验工具', '归类依据']
    widths = [15, 8, 10, 24, 15, 8, 16, 30, 16, 46, 12, 13, 10]

    used = set()
    project_order = [p[0] for p in projects]
    for pname, _, files in projects:
        ws = wb.create_sheet(safe_sheet_name(pname, used))
        ws.sheet_view.showGridLines = False
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.merge_cells('A1:M1')
        ws['A1'] = f'{pname} · SIP 方法论顺序对应表'
        ws['A1'].font = Font(name='微软雅黑', size=14, bold=True, color='1F3864')
        ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
        ws.row_dimensions[1].height = 28
        ws.merge_cells('A2:M2')
        ws['A2'] = ('主顺序＝SIP 九个环节（1→9）；子顺序＝每个环节内部执行次序（1级优先，12级最后）。'
                    '原表标黄内容不纳入。归类依据列：标准答案/同类工序/规则推断。')
        ws['A2'].font = Font(name='微软雅黑', size=9, color='808080')
        ws['A2'].alignment = Alignment(vertical='center', wrap_text=True)
        ws.row_dimensions[2].height = 30
        for i, h in enumerate(headers, 1):
            c = ws.cell(row=3, column=i, value=h)
            c.font = head_font
            c.fill = head_fill
            c.alignment = center
            c.border = border

        file_idx = {os.path.splitext(f)[0]: i for i, f in enumerate(files)}
        rows = [x for x in all_rows if x['project'] == pname]
        rows.sort(key=lambda x: (MAIN_IDX[x['main']], SUB_IDX[x['sublevel']],
                                 file_idx.get(x['file'], 999), x['src']))
        r = 4
        for main in MAIN_ORDER:
            grp = [x for x in rows if x['main'] == main]
            if not grp:
                continue
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=13)
            c = ws.cell(row=r, column=1, value=f'▶ {main}　（{len(grp)} 项）')
            c.font = main_font
            c.fill = main_fill
            c.alignment = Alignment(vertical='center')
            for col in range(1, 14):
                ws.cell(row=r, column=col).fill = main_fill
                ws.cell(row=r, column=col).border = border
            ws.row_dimensions[r].height = 22
            r += 1
            for sub in SUB_ORDER:
                sgrp = [x for x in grp if x['sublevel'] == sub]
                if not sgrp:
                    continue
                ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=13)
                c = ws.cell(row=r, column=1, value=f'　　{sub}　（{len(sgrp)} 项）')
                c.font = sub_font
                c.alignment = Alignment(vertical='center')
                for col in range(1, 14):
                    ws.cell(row=r, column=col).fill = sub_fill
                    ws.cell(row=r, column=col).border = border
                r += 1
                for x in sgrp:
                    vals = [x['main'], x['sublevel'], x['file'], x['product'],
                            x['model'], x['seq'], x['stage'], x['item'], x['sub'],
                            x['std'], x['method'], x['tool'], x['source']]
                    for i, v in enumerate(vals, 1):
                        c = ws.cell(row=r, column=i, value=v)
                        c.font = body_font
                        c.border = border
                        c.alignment = center if i in (1, 2, 3, 5, 6, 11, 12, 13) else wrap
                    r += 1
        ws.freeze_panes = 'C4'
        ws.auto_filter.ref = f'A3:M{r - 1}'

    # --------------------------- 说明
    ws = wb.create_sheet('说明', 0)
    ws.sheet_view.showGridLines = False
    ws.column_dimensions['A'].width = 20
    ws.column_dimensions['B'].width = 16
    ws.column_dimensions['C'].width = 78
    ws.merge_cells('A1:C1')
    ws['A1'] = 'SIP 方法论顺序对应表 · 说明'
    ws['A1'].font = Font(name='微软雅黑', size=14, bold=True, color='1F3864')
    r = 3

    verify_text = '未提供标准答案文件'
    if ref:
        if verify:
            parts = []
            for f, ok, bad, _ in verify:
                parts.append(f'{f}：{ok}/{ok + bad} 行一致' + ('，0 行不符' if bad == 0 else f'，{bad} 行不符'))
            verify_text = '；'.join(parts) if parts else '标准答案已加载，但输入中没有匹配的机型文件'
        else:
            verify_text = '标准答案已加载（未找到匹配的机型文件，未做逐行比对）'

    intro = [
        ('任务', '把《成品检验指导书》中的每条检验，按 SIP 方法论两层顺序模型归类：'
                 '主顺序（9 个环节）+ 子顺序（环节内部 1级→12级）。'),
        ('标准答案校验', verify_text),
        ('标黄处理', '原表标黄的“检验频率、重要等级、检验记录表、备注”四列及黄色文件编号/版本不纳入输出；'
                   '浅蓝色版本更新标注保留。'),
        ('共用项处理', '团队标准答案没有单独的“共用”区：同一工序中既涉及水又涉及气的行，'
                    '按行内容拆分到“4.通气检验”或“5.通水检验”。'),
        ('归类依据', '每条记录最后一列：标准答案＝标准答案中已有同名工序/项目；'
                  '同类工序＝按同一语义套用的新机型工序；规则推断＝关键字兜底，需人工复核。'),
        ('标黄跳过行', f'解析时跳过序号单元格为黄色的行：{skipped_yellow} 行。'),
    ]
    for k, v in intro:
        ws.cell(row=r, column=1, value=k).font = Font(name='微软雅黑', size=10, bold=True)
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        c = ws.cell(row=r, column=2, value=v)
        c.font = Font(name='微软雅黑', size=10)
        c.alignment = Alignment(vertical='top', wrap_text=True)
        ws.row_dimensions[r].height = 15 * (len(v) // 70 + 1) + 8
        r += 1

    r += 1
    ws.cell(row=r, column=1, value='两层顺序模型').font = Font(name='微软雅黑', size=11, bold=True)
    r += 1
    for i, h in enumerate(['主顺序', '子顺序', '该级的含义'], 1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = head_font
        c.fill = head_fill
        c.alignment = center
        c.border = border
    r += 1
    for main in MAIN_ORDER:
        for sub, mean in SUB_MEANING[main]:
            ws.cell(row=r, column=1, value=main).font = body_font
            ws.cell(row=r, column=2, value=sub).font = body_font
            c = ws.cell(row=r, column=3, value=mean)
            c.font = body_font
            c.alignment = Alignment(vertical='center', wrap_text=True)
            for col in range(1, 4):
                ws.cell(row=r, column=col).border = border
            r += 1

    r += 1
    ws.cell(row=r, column=1, value='数据源清单').font = Font(name='微软雅黑', size=11, bold=True)
    r += 1
    for i, h in enumerate(['项目', '机型文件', '产品名称 / 机型代码'], 1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = head_font
        c.fill = head_fill
        c.alignment = center
        c.border = border
    r += 1
    for pname, _, files in projects:
        for f in files:
            stem = os.path.splitext(f)[0]
            row = next((x for x in all_rows if x['project'] == pname and x['file'] == stem), None)
            ws.cell(row=r, column=1, value=pname).font = body_font
            ws.cell(row=r, column=2, value=stem).font = body_font
            ws.cell(row=r, column=3,
                    value=(f"{row['product']} / {row['model']}" if row else '')).font = body_font
            for col in range(1, 4):
                ws.cell(row=r, column=col).border = border
            r += 1

    r += 1
    ws.cell(row=r, column=1, value='环节统计').font = Font(name='微软雅黑', size=11, bold=True)
    r += 1
    stat_head = r
    hdr = ['主顺序'] + [p[0] for p in projects]
    for i, h in enumerate(hdr, 1):
        c = ws.cell(row=r, column=i, value=h)
        c.font = head_font
        c.fill = head_fill
        c.alignment = center
        c.border = border
    r += 1
    for main in MAIN_ORDER:
        ws.cell(row=r, column=1, value=main).font = body_font
        ws.cell(row=r, column=1).border = border
        for i, (pname, _, _) in enumerate(projects, 2):
            ws.cell(row=r, column=i, value=sum(
                1 for x in all_rows if x['project'] == pname and x['main'] == main))
            ws.cell(row=r, column=i).font = body_font
            ws.cell(row=r, column=i).border = border
            ws.cell(row=r, column=i).alignment = center
        r += 1
    ws.cell(row=r, column=1, value='合计').font = Font(name='微软雅黑', size=10, bold=True)
    ws.cell(row=r, column=1).border = border
    for i, (pname, _, _) in enumerate(projects, 2):
        c = ws.cell(row=r, column=i,
                    value=sum(1 for x in all_rows if x['project'] == pname))
        c.font = Font(name='微软雅黑', size=10, bold=True)
        c.border = border
        c.alignment = center

    # 同类工序统计
    if ref:
        same = {}
        for x in all_rows:
            if x['source'] == '同类工序':
                key = (x['stage'], x['main'], x['sublevel'])
                same[key] = same.get(key, 0) + 1
        if same:
            r += 2
            ws.cell(row=r, column=1,
                    value='同类工序统计（建议抽查）').font = Font(name='微软雅黑', size=11, bold=True)
            r += 1
            for i, h in enumerate(['检验工序', '主顺序', '子顺序'], 1):
                c = ws.cell(row=r, column=i, value=h)
                c.font = head_font
                c.fill = head_fill
                c.alignment = center
                c.border = border
            r += 1
            for (stage, main, sub), cnt in sorted(
                    same.items(), key=lambda kv: (MAIN_IDX[kv[0][1]], SUB_IDX[kv[0][2]], kv[0][0])):
                ws.cell(row=r, column=1, value=stage).font = body_font
                ws.cell(row=r, column=2, value=main).font = body_font
                ws.cell(row=r, column=3, value=f'{sub}（{cnt} 项）').font = body_font
                for col in range(1, 4):
                    ws.cell(row=r, column=col).border = border
                r += 1

    # --------------------------- 待确认项
    ws2 = wb.create_sheet('待确认项（规则推断）')
    ws2.sheet_view.showGridLines = False
    for col, w in zip('ABCDE', [10, 22, 32, 16, 10]):
        ws2.column_dimensions[col].width = w
    ws2.merge_cells('A1:E1')
    ws2['A1'] = '规则推断项清单（新机型特有工序，建议文件编写同事复核）'
    ws2['A1'].font = Font(name='微软雅黑', size=12, bold=True, color='C00000')
    for i, h in enumerate(['机型文件', '检验工序', '检验项目', '主顺序', '子顺序'], 1):
        c = ws2.cell(row=3, column=i, value=h)
        c.font = head_font
        c.fill = head_fill
        c.alignment = center
        c.border = border
    pending = [x for x in all_rows if x['source'] == '规则推断']
    if not pending:
        ws2.cell(row=4, column=1, value='无（所有工序均按标准答案或同类工序规则归类）').font = body_font
    else:
        rr = 4
        for x in pending:
            for i, v in enumerate([x['file'], x['stage'], x['item'],
                                   x['main'], x['sublevel']], 1):
                c = ws2.cell(row=rr, column=i, value=v)
                c.font = body_font
                c.border = border
            rr += 1

    os.makedirs(os.path.dirname(os.path.abspath(out_path)) or '.', exist_ok=True)
    wb.save(out_path)


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser(description='按 SIP 方法论归类成品检验指导书')
    ap.add_argument('--input', required=True, help='zip 或包含 xlsx 的文件夹')
    ap.add_argument('--output', required=True, help='输出 xlsx 路径')
    ap.add_argument('--reference', help='标准答案 xlsx（含 主顺序/子顺序 列）')
    ap.add_argument('--project', help='当 xlsx 直接放在根目录时，指定项目名')
    ap.add_argument('--keep-temp', action='store_true', help='保留 zip 解压临时目录')
    args = ap.parse_args()

    tmp = None
    root = args.input
    if os.path.isfile(root) and root.lower().endswith('.zip'):
        tmp = tempfile.mkdtemp(prefix='sip_input_')
        print(f'解压 zip → {tmp}')
        extract_zip(root, tmp)
        root = tmp
    if not os.path.isdir(root):
        raise SystemExit(f'输入不存在或不是目录/zip: {args.input}')

    ref = load_reference(args.reference) if args.reference else None
    projects = discover_projects(root, args.project)
    print('项目：')
    for pname, pdir, files in projects:
        print(f'  {pname}: {len(files)} 个文件 ({pdir})')

    all_rows = []
    skipped_total = 0
    for pname, pdir, files in projects:
        for f in files:
            path = os.path.join(pdir, f)
            rows, skipped = parse_sheet(path, pname)
            skipped_total += skipped
            for row in rows:
                (main, sub), source = classify(
                    row['stage'], row['item'], row['sub'], row['std'], row['method'],
                    ref['pairs'] if ref else None,
                    ref['stages'] if ref else None)
                row['main'] = main
                row['sublevel'] = sub
                row['source'] = source
            all_rows.extend(rows)
            print(f'  {f}: {len(rows)} 行' + (f'（跳过黄行 {skipped}）' if skipped else ''))

    if not all_rows:
        raise SystemExit('没有解析到任何有效数据行')

    verify = verify_reference(all_rows, ref)
    if verify:
        print('标准答案校验：')
        for f, ok, bad, diffs in verify:
            print(f'  {f}: {ok}/{ok + bad} 一致, {bad} 不符')
            for d in diffs:
                print('    ', d)

    write_workbook(all_rows, projects, args.output, ref, verify, skipped_total)
    print(f'\n已生成: {os.path.abspath(args.output)}')
    print(f'总行数: {len(all_rows)}')
    from collections import Counter
    cnt = Counter(x['main'] for x in all_rows)
    for main in MAIN_ORDER:
        if cnt[main]:
            print(f'  {main}: {cnt[main]}')
    src = Counter(x['source'] for x in all_rows)
    print('归类依据:', dict(src))
    pending = src.get('规则推断', 0)
    if pending:
        print(f'⚠ 有 {pending} 行属于规则推断，见“待确认项”工作表')

    if tmp and not args.keep_temp:
        shutil.rmtree(tmp, ignore_errors=True)
    return 0


if __name__ == '__main__':
    sys.exit(main())
