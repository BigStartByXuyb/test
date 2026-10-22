# -*- coding: utf-8 -*-
"""SIP 方法论归类规则（机器可执行口径）。

主顺序：9 个环节；子顺序：每个环节内部 1级→12级。
规则支持用任意标准答案 xlsx 逐行校验（由调用方通过 load_reference 传入，不假定文件名或路径）。
新增机型特有工序时，优先补 STAGE_RULES / EXACT / 行级拆分，不要改动标准答案已有归属。
"""
import re

M = '1. 铭牌核对'
W = '2. 外观检验'
D = '3. 通电检验'
Q = '4. 通气检验'
S = '5. 通水检验'
L = '6. 老化测试'
G = '7. 工艺验证'
U = '8. 通断电检验'
B = '9. 备份杀毒检验'

MAIN_ORDER = [M, W, D, Q, S, L, G, U, B]
MAIN_IDX = {m: i for i, m in enumerate(MAIN_ORDER)}
SUB_ORDER = [f'{i}级' for i in range(1, 13)]
SUB_IDX = {s: i for i, s in enumerate(SUB_ORDER)}


def norm(s):
    """归一化工序/项目文本：去换行、空格、全角括号。"""
    if s is None:
        return ''
    return (str(s).replace('\n', '').replace('\r', '').replace(' ', '')
            .replace('　', '').replace('（', '(').replace('）', ')'))


HEADER_BASES = (
    '序号', '检验工序', '检验项目', '检验标准', '检验方法', '检验工具',
    '主顺序', '子顺序', '检验频率', '重要等级', '检验记录表', '备注',
)


def header_key(v):
    """把“子顺序（1级优先进行…）”这类带说明的表头归一成基础列名。"""
    v = norm(v)
    for base in HEADER_BASES:
        if v == base or v.startswith(base):
            return base
    return v


# ---------------------------------------------------------------- 工序级规则
STAGE_RULES = {
    '机台核对': (M, '1级'), 'Y轴精度检验': (D, '2级'), '精度检查表': (D, '4级'),
    '厂务动力检查': (D, '1级'), '补偿文件': (D, '3级'), '精度补偿': (D, '2级'),
    '测高': (D, '5级'), 'NCS测高行程': (D, '5级'), '升降台水平': (D, '5级'),
    '规整': (D, '5级'), '搬运手臂': (D, '5级'), '下手臂吸盘切换': (D, '5级'),
    '清洗': (D, '5级'), 'UV': (D, '5级'), '全自动空切流程': (D, '5级'),
    '清洗桶水帘': (D, '5级'), 'BBD': (D, '5级'),
    '视觉对位及标定': (D, '7级'), '相机验证': (D, '7级'), '相机对位': (D, '7级'),
    'X、Y轴垂直度确认': (D, '7级'), '相机': (D, '7级'),
    '气缸示教': (D, '8级'),
    '接地电阻': (D, '9级'), '接地': (D, '9级'), '漏电电压': (D, '9级'),
    '主轴转速': (D, '10级'), '设备噪音': (D, '10级'),
    '主轴与安全门联锁': (D, '11级'), '设备序列号': (D, '11级'),
    '安全防护检查': (D, '12级'), '版本控制': (D, '12级'), '条码枪': (D, '12级'),
    '法兰与刀片': (D, '12级'),
    '空气过滤组合': (Q, '1级'), '空气过滤组合(切割机)': (Q, '1级'),
    '空气过滤组合(分选机)': (Q, '1级'),
    '气枪': (Q, '2级'), '高效气枪': (Q, '2级'),
    '真空发生器及电磁阀': (Q, '3级'), '冷扩腔室密封性': (Q, '3级'),
    '冷扩腔室入口气幕流量': (Q, '2级'), '吹气杆吹气': (Q, '2级'),
    '漏水检测': (S, '3级'), '白色漏水绳': (S, '3级'), '发泡机测试': (S, '3级'),
    '防水及防锈处理': (S, '3级'), '切割门钣金漏水检': (S, '3级'),
    '水雾检测': (S, '3级'), 'BW366防锈剂涂抹确认': (S, '3级'),
    '清洗过程及效果': (S, '1级'),
    '老化测试': (L, '2级'), '老化前确认事项': (L, '1级'),
    '切割精度验证': (G, '1级'), '切割品质验证': (G, '2级'),
    '刀痕宽度': (G, '3级'), '工艺验证UV膜刀痕检验': (G, '3级'),
    '工艺验证Dummy切深检验': (G, '3级'), '工艺验收': (G, '1级'),
    '工艺验收标准': (G, '1级'),
    '断电确认': (U, '1级'),
    'Windows10自动更新关闭': (B, '2级'), 'windows激活': (B, '2级'),
    '机台备份': (B, '3级'), '软件确认': (B, '4级'),
    '杀毒软件': (B, '5级'), 'USB确认': (B, '8级'),
    # 外观类
    '水槽盖板': (W, '1级'), '压力表': (W, '1级'), '地脚': (W, '1级'),
    '地脚检验': (W, '1级'), '接头': (W, '1级'), '设备颜色': (W, '1级'),
    '压铁环气缸连接件顶丝确认': (W, '1级'),
    '手动操作': (W, '2级'), '切割门': (W, '2级'),
    '主轴冷却水接口': (W, '3级'), '主轴轴心': (W, '3级'),
    '屏幕后盖板': (W, '3级'),
    '外观': (W, '1级'), '机台整理': (W, '2级'), '安全注意及防护事项': (W, '1级'),
    '整体布局': (W, '2级'), '管道及量具标识': (W, '2级'),
    '机台外罩检验': (W, '1级'), '配件检验': (M, '2级'),
    '清洁检验': (W, '1级'), '出货固定确认': (W, '1级'),
    '出货称重确认': (W, '1级'), '显示屏侧物料防护': (W, '1级'),
    'X轴防尘罩': (W, '3级'), 'NCS固定钣金': (W, '3级'),
    '拖链线': (W, '2级'), '网线': (W, '2级'), '运动部件线材': (W, '2级'),
    '总电源开关': (W, '1级'), '强弱电分离': (W, '2级'),
    '热扩台盘调节机构': (W, '3级'),
    '切割区门板': (W, '1级'), '光纤检验': (W, '1级'),
    '水气路板螺纹': (W, '3级'), '主轴冷却水和供气管': (W, '3级'),
    # 通电/功能类
    '料库复位': (D, '5级'), '规整复位': (D, '5级'), '搬运复位': (D, '5级'),
    '清洗组复位': (D, '5级'), '切割组复位': (D, '5级'),
    '横梁组搬运复位': (D, '5级'), '单臂机械手复位': (D, '5级'),
    '上料搬运组复位': (D, '5级'), '下料搬运组复位': (D, '5级'),
    '清洗搬运组复位': (D, '5级'),
    '料框上料(第一工位)': (D, '5级'), '规整调整(第二工位)': (D, '5级'),
    '单臂机械手搬片(第二工位)': (D, '5级'), '规整(第三工位)': (D, '5级'),
    '上料抓取机械手(Z向)(第三工位)': (D, '5级'),
    '上料抓取机械手(Z向)第三工位': (D, '5级'),
    '上料搬运(第四工位)': (D, '5级'),
    '视觉对位及标定(第四工位)': (D, '7级'),
    '视觉对位及标定(第五工位)': (D, '7级'),
    '测高(第五工位)': (D, '5级'), '测高(第六工位)': (D, '5级'),
    '切割(第六工位)': (D, '5级'), '切割(第七工位)': (D, '5级'),
    '切割': (D, '5级'),
    '搬运手臂(第七工位)': (D, '5级'), '搬运手臂(第九工位)': (D, '5级'),
    '清洗组件(第八工位)': (D, '5级'), '下料搬运(第八工位)': (D, '5级'),
    '背洗(第九工位)': (D, '5级'), '清洗搬运(第十工位)': (D, '5级'),
    '清洗(第十一工位)': (D, '5级'), '下料(第十二工位)': (D, '5级'),
    'I/O信号点位确认': (D, '1级'), '电子断路器测试标准': (D, '9级'),
    '料库单元': (D, '5级'), '上料单元': (D, '5级'), '下料单元': (D, '5级'),
    '清洗单元': (D, '5级'), '干燥单元': (D, '5级'), '转盘单元': (D, '5级'),
    '吸嘴单元': (D, '5级'), '基板吸取器': (D, '5级'),
    '单元吸取器': (D, '5级'), '托盘单元': (D, '5级'),
    '接触及非接触测高': (D, '5级'), '排风电机': (D, '1级'),
    '整机': (D, '12级'), 'EFEM机构水平': (D, '5级'),
    '预校准工作台': (D, '5级'), '预定位台': (D, '5级'), 'LSD': (D, '5级'),
    '接地检测': (D, '9级'), 'Foup开盖': (D, '5级'), '功能验收': (D, '5级'),
    '搬送手臂': (D, '5级'), '冷扩工作动作': (D, '5级'),
    '热扩工作动作': (D, '5级'), '热风枪加热功能': (D, '5级'),
    '热扩盘加热功能': (D, '5级'), '红外加热功能': (D, '5级'),
    'UV功能': (D, '5级'), '升降开关门': (D, '5级'),
    '全自动传片测试': (D, '5级'), '首料检功能': (D, '5级'),
    '固定式扫码器': (D, '12级'), '限位片检验': (D, '11级'),
    '机台外壳': (D, '9级'), '陶瓷手臂': (D, '9级'),
    '热扩台': (D, '9级'), '热扩升降平台': (D, '9级'),
    '左侧门板玻璃': (D, '9级'), '右侧门板玻璃': (D, '9级'),
    '安全门板玻璃': (D, '9级'),
    '非金属部件(坦克链，气管，线槽)': (D, '9级'), '磁环': (D, '9级'),
    '横移臂防尘板': (D, '9级'), '旋转臂防尘板': (D, '9级'),
    '离子风扇': (D, '9级'),
}
for _n in ['上料吸嘴1', '上料吸嘴2', '上料吸嘴3', '上料吸嘴4',
           '下料吸嘴1', '下料吸嘴2', '下料吸嘴3', '下料吸嘴4',
           '横移臂吸嘴1', '横移臂吸嘴2', '横移臂吸嘴3', '横移臂吸嘴4',
           '旋转臂吸嘴1', '旋转臂吸嘴2', '旋转臂吸嘴3', '旋转臂吸嘴4',
           '上料防尘盖板', '下料防尘盖板', '切割台', '清洗台',
           '料库门板玻璃', '切割门板玻璃', '分选门板玻璃',
           '上料防静电橡胶顶丝1-3', '上料防静电橡胶顶丝4-6',
           '上料防静电橡胶顶丝7-9',
           '下料防静电橡胶顶丝1-3', '下料防静电橡胶顶丝4-6',
           '下料防静电橡胶顶丝7-9']:
    STAGE_RULES[_n] = (D, '9级')

# ---------------------------------------------------------------- 项目级例外
EXACT = {
    ('机台核对', '受订单核对'): (M, '1级'),
    ('机台核对', '履历表完整性'): (W, '1级'),
    ('Y轴精度检验', '重复精度'): (D, '2级'),
    ('Y轴精度检验', '定位精度检验'): (D, '2级'),
    ('流量、压力及位置检查表', '相机焦距高度'): (D, '5级'),
    ('流量、压力及位置检查表', 'Y1-Y2最小轴间距位置确认'): (D, '5级'),
    ('流量、压力及位置检查表', '接触式测高位置及刀刃外观确认'): (D, '5级'),
    ('流量、压力及位置检查表', '喷水架位置检查'): (S, '2级'),
    ('流量、压力及位置检查表', '工作台夹爪状态'): (Q, '2级'),
    ('流量、压力及位置检查表', '清洗二流体气压'): (S, '1级'),
    ('流量、压力及位置检查表', 'BBD 传感器'): (S, '1级'),
    ('流量、压力及位置检查表', 'NCS 传感器'): (S, '1级'),
    ('流量、压力及位置检查表', '切割水温度传感器(Z1/Z2 Blade温度)'): (S, '1级'),
    ('流量、压力及位置检查表', '主轴温度传感器(Z1/Z2 主轴冷却水温度)'): (S, '1级'),
    ('流量、压力及位置检查表', '环境温度传感器'): (S, '1级'),
    ('流量、压力及位置检查表', '水帘角度'): (S, '1级'),
    ('流量、压力及位置检查表', '气帘角度'): (S, '1级'),
    ('流量、压力及位置检查表', '水气帘高度'): (S, '1级'),
    ('流量、压力及位置检查表', '切割室隔板高度'): (S, '1级'),
    ('加密狗', '核对规格'): (M, '2级'),
    ('加密狗', '永久版本'): (B, '1级'),
    ('加密狗', '序列号对应'): (B, '1级'),
    ('工控机', '刀轮工控机'): (M, '2级'),
    ('工控机', '通电后核对系统配置'): (D, '12级'),
    ('工控机', '确认断电后是否对机台造成影响'): (U, '2级'),
    ('杀毒软件', '《火绒杀毒软件》是否安装'): (B, '5级'),
    ('杀毒软件', '查看查杀病毒记录'): (B, '6级'),
    ('杀毒软件', '查看加密有无执行'): (B, '7级'),
}

BOM_STAGES = {
    'UPS电源', '交换机', '继电器', '触摸屏', '触控屏', '显示器',
    '报警灯', 'LED防水灯', 'LED灯', '急停', '气缸', '主轴', '风机',
    '控制系统', '电阻',
}
BOM_RE = re.compile(r'^\s*\d+\s*pcs', re.I)


# ---------------------------------------------------------------- 分类
def _classify_raw(stage, item, sub, std, method):
    key = (stage, item)
    if key in EXACT:
        if item == '接触式测高位置及刀刃外观确认':
            return (D, '6级') if '刀刃外观' in std else (D, '5级')
        return EXACT[key]
    if stage.startswith('流量、压力及位置检查'):
        if 'NCS检查' in item:
            return (S, '1级') if ('水柱' in std or '冲洗' in std) else (Q, '2级')
        if '水气帘' in item:
            return (S, '1级')
        if '气帘角度' in item or '水帘角度' in item:
            return (S, '1级')
        if '吹气' in item or '光尺气帘' in item:
            return (Q, '2级')
        if '真空' in item:
            return (Q, '3级')
        if '二流体' in item:
            return (S, '1级')
        if '喷水架' in item:
            return (S, '2级')
        if any(k in item for k in ('自清洁', '水帘', '水压', '清洗水')) or '水柱' in std:
            return (S, '1级')
        if any(k in item for k in ('气帘', '正压', '压力', '氮气', '热风', '红外')):
            return (Q, '2级')
        if any(k in item for k in ('气缸', '位置', '水平', '轴间距', '测高行程')):
            return (D, '5级')
        if any(k in item for k in ('流量', '传感', '高度', '隔板')):
            return (S, '1级')
    if stage == '工控机':
        text = item + std
        if '断电' in text:
            return (U, '2级')
        if ('序列号' in text or 'S/N' in text or '核对规格' in item
                or BOM_RE.match(std)):
            return (M, '2级')
        return (D, '12级')
    if stage == '料库检料传感器确认':
        return (D, '6级')
    if stage == '气管、水管':
        return (S, '3级') if ('水管' in std and '气管' not in std) else (Q, '4级')
    if stage == '离子棒':
        return (W, '1级') if ('安装位置' in std or '安装在' in std) else (D, '9级')
    if stage == 'NCS光纤':
        return (W, '1级') if ('外观' in std or '无挤压' in std) else (W, '2级')
    if stage == '安全防护检查':
        return (S, '3级') if ('漏水' in std or '漏水' in item) else (D, '12级')
    if stage == '陶瓷手臂':
        return (W, '1级') if '拆解打包' in (item + std) else (D, '9级')
    if stage in ('工艺验收', '工艺验收标准'):
        if any(k in (item + std) for k in ('尺寸', '厚度', '宽度', '深度')):
            return (G, '1级')
        if any(k in (item + std) for k in ('外观', '崩边', '沾污', '破裂', '分段', '分离')):
            return (G, '2级')
        return (G, '1级')
    if stage in STAGE_RULES:
        return STAGE_RULES[stage]
    if stage in BOM_STAGES:
        if BOM_RE.match(std):
            return (M, '2级')
        if any(k in std for k in ('外观', '安装', '完好', '牢固', '无损伤')):
            return (W, '1级')
        return (D, '5级')
    # 关键字推断；返回 None 表示完全无法推断
    if any(k in stage for k in ('精度', '补偿')):
        return (D, '4级')
    if any(k in stage for k in ('外观', '整理', '清洁', '标识', '布局', '门板', '盖板',
                                '接头', '地脚', '颜色', '外壳')):
        return (W, '1级')
    if any(k in stage for k in ('复位', '气缸', '电机', '转速', '噪音', '功能', '流程',
                                '测试', '点检', '单元', '工位')):
        return (D, '5级')
    if any(k in stage for k in ('真空', '气密', '气幕')) or (
            '气' in stage and '水' not in stage):
        return (Q, '2级')
    if '水' in stage or '漏' in stage:
        return (S, '3级')
    if '老化' in stage:
        return (L, '2级')
    if '工艺' in stage:
        return (G, '1级')
    if '断电' in stage:
        return (U, '1级')
    if any(k in stage for k in ('备份', '杀毒', '病毒', 'USB', 'window', '加密')):
        return (B, '3级')
    return None


def _explicit_match(stage, item):
    """是否命中了明确规则（而不是关键字推断）。"""
    if (stage, item) in EXACT:
        return True
    if stage.startswith('流量、压力及位置检查'):
        return True
    if stage in ('工控机', '料库检料传感器确认', '气管、水管', '离子棒', 'NCS光纤',
                 '安全防护检查', '陶瓷手臂', '工艺验收', '工艺验收标准'):
        return True
    return stage in STAGE_RULES or stage in BOM_STAGES


def classify(stage, item, sub='', std='', method='',
             ref_pairs=None, ref_stages=None):
    """返回 ((主顺序, 子顺序), 归类依据)。

    归类依据：
      标准答案  - 标准答案中存在同名 (工序, 项目)
      同类工序  - 标准答案中存在同名工序，或命中了明确的工序/项目规则
      规则推断  - 关键字推断（只会落在主顺序 2–9、子顺序 1–5 级；完全无法推断时用 (3. 通电检验, 5级)），需人工复核
    """
    result = _classify_raw(stage, item, sub, std, method)
    if result is None:
        return (D, '5级'), '规则推断'
    if ref_pairs is not None and (stage, item) in ref_pairs:
        return result, '标准答案'
    if ref_stages is not None and stage in ref_stages:
        return result, '同类工序'
    if _explicit_match(stage, item):
        return result, '同类工序'
    return result, '规则推断'


# ---------------------------------------------------------------- 标准答案读取
def load_reference(path):
    """读取同事确认的标准答案 xlsx。

    返回 dict：
      pairs: {(工序, 项目)}
      stages: {工序}
      seq: {原序号: (主顺序, 子顺序)}
      sheet_title / product / model
    """
    from openpyxl import load_workbook
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    header_row = None
    for r in range(1, min(ws.max_row, 12) + 1):
        vals = [header_key(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if ('序号' in vals and '检验工序' in vals
                and '主顺序' in vals and '子顺序' in vals):
            header_row = r
            break
    if header_row is None:
        raise ValueError('标准答案中找不到包含 序号/检验工序/主顺序/子顺序 的表头行')

    col = {}
    for c in range(1, ws.max_column + 1):
        v = header_key(ws.cell(header_row, c).value)
        if v and v not in col:
            col[v] = c
    missing = [x for x in ('序号', '检验工序', '检验项目', '主顺序', '子顺序')
               if x not in col]
    if missing:
        raise ValueError(
            '标准答案缺少必需列: ' + '、'.join(missing)
            + '（第一个工作表的表头需含 序号/检验工序/检验项目/主顺序/子顺序）')
    c_seq = col['序号']
    c_stage = col['检验工序']
    c_item = col['检验项目']
    c_main = col['主顺序']
    c_sub = col['子顺序']

    pairs, stages, seq = set(), set(), {}
    product = model = ''
    # 只扫描表头之前的 10 行，不含表头及其下方数据行
    for r in range(max(1, header_row - 10), header_row):
        for c in range(1, ws.max_column + 1):
            v = ws.cell(r, c).value
            if not isinstance(v, str):
                continue
            if v.startswith('产品名称'):
                product = v.split('：', 1)[-1].strip()
            elif v.startswith('机型代码'):
                model = v.split('：', 1)[-1].strip()

    st = it = ''
    for r in range(header_row + 1, ws.max_row + 1):
        a = ws.cell(r, c_seq).value
        if isinstance(a, (int, float)):
            seq_no = int(a)
        elif isinstance(a, str) and a.strip().isdigit():
            seq_no = int(a.strip())
        else:
            continue
        b = ws.cell(r, c_stage).value
        if b is not None and str(b).strip():
            st = norm(b)
            it = ''
        c = ws.cell(r, c_item).value
        if c is not None and str(c).strip():
            it = norm(c)
        main = ws.cell(r, c_main).value
        sub = ws.cell(r, c_sub).value
        if main and sub:
            pairs.add((st, it))
            stages.add(st)
            seq[seq_no] = (str(main).strip(), str(sub).strip())
    return {
        'pairs': pairs,
        'stages': stages,
        'seq': seq,
        'sheet_title': ws.title,
        'product': product,
        'model': model,
    }
