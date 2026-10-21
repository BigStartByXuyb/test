# -*- coding: utf-8 -*-
"""源表解析与项目发现（同格式生成与分析视图共用）。"""
import os

from openpyxl import load_workbook

from sip_rules import norm, header_key


def is_yellow(cell):
    f = cell.fill
    if f is None or f.fill_type is None:
        return False
    c = f.start_color
    rgb = getattr(c, 'rgb', None)
    if isinstance(rgb, str) and len(rgb) == 8:
        r, g, b = int(rgb[2:4], 16), int(rgb[4:6], 16), int(rgb[6:8], 16)
        return r >= 200 and g >= 180 and b <= 160
    return False


def cell_text(ws, row, col):
    if col is None:
        return ''
    cell = ws.cell(row, col)
    if is_yellow(cell):
        return ''
    v = cell.value
    return str(v).strip() if v is not None else ''


def find_header_row(ws):
    for r in range(1, min(ws.max_row, 15) + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if '序号' in vals and '检验工序' in vals and '检验项目' in vals:
            return r
    raise ValueError(f'{ws.title}: 找不到表头行（需要 序号/检验工序/检验项目）')


def get_columns(ws, header_row):
    cmap = {}
    for c in range(1, ws.max_column + 1):
        v = header_key(ws.cell(header_row, c).value)
        if v and v not in cmap:
            cmap[v] = c
    missing = [x for x in ('序号', '检验工序', '检验项目') if x not in cmap]
    if missing:
        raise ValueError(f'{ws.title}: 缺少必需列 {missing}')
    return cmap


def parse_sheet(path, project):
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    hr = find_header_row(ws)
    cmap = get_columns(ws, hr)
    c_seq = cmap['序号']
    c_stage = cmap['检验工序']
    c_item = cmap['检验项目']
    c_std = cmap.get('检验标准')
    c_method = cmap.get('检验方法')
    c_tool = cmap.get('检验工具')

    sub_cols = []
    if c_std:
        sub_cols = [c for c in range(c_item + 1, c_std) if c != c_stage]

    product = model = ''
    for r in range(1, hr):
        for c in range(1, ws.max_column + 1):
            v = cell_text(ws, r, c)
            if v.startswith('产品名称'):
                product = v.split('：', 1)[-1].strip()
            elif v.startswith('机型代码'):
                model = v.split('：', 1)[-1].strip()

    rows = []
    stage = item = ''
    skipped_yellow = 0
    for r in range(hr + 1, ws.max_row + 1):
        seqv = ws.cell(r, c_seq).value
        if isinstance(seqv, (int, float)):
            seq = int(seqv)
        elif isinstance(seqv, str) and seqv.strip().isdigit():
            seq = int(seqv.strip())
        else:
            continue
        if is_yellow(ws.cell(r, c_seq)):
            skipped_yellow += 1
            continue
        b = cell_text(ws, r, c_stage)
        if b:
            stage = norm(b)
            item = ''
        c = cell_text(ws, r, c_item)
        if c:
            item = norm(c)
        sub = ' / '.join(x for x in (cell_text(ws, r, sc) for sc in sub_cols) if x)
        rows.append(dict(
            project=project,
            file=os.path.splitext(os.path.basename(path))[0],
            product=product,
            model=model,
            seq=seq,
            stage=stage,
            item=item,
            sub=sub,
            std=cell_text(ws, r, c_std),
            method=cell_text(ws, r, c_method),
            tool=cell_text(ws, r, c_tool),
            src=r,
        ))
    return rows, skipped_yellow


def _xlsx_files(d):
    return [f for f in sorted(os.listdir(d))
            if f.lower().endswith('.xlsx') and not f.startswith('~$')]


def _data_root(d):
    if _xlsx_files(d):
        return d
    subs = [x for x in sorted(os.listdir(d)) if os.path.isdir(os.path.join(d, x))]
    if len(subs) == 1:
        return _data_root(os.path.join(d, subs[0]))
    return d


def discover_projects(root, project_name=None):
    root = _data_root(root)
    projects = []
    for entry in sorted(os.listdir(root)):
        p = os.path.join(root, entry)
        if os.path.isdir(p):
            files = _xlsx_files(p)
            if files:
                projects.append((entry, p, files))
    root_files = _xlsx_files(root)
    if root_files:
        name = project_name or os.path.basename(os.path.abspath(root)) or '默认项目'
        projects.append((name, root, root_files))
    if not projects:
        raise ValueError(f'在 {root} 下没有找到 .xlsx 文件')
    return projects
