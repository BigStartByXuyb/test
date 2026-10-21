# -*- coding: utf-8 -*-
"""按参考样张格式生成对应表（每个机型一个 Sheet）。

与参考样张一致：
  - 删除源表顶部标题行与页脚，表头位于第 1 行
  - 保留原 A-L 列与合并单元格（整体上移）
  - 新增 M 主顺序 / N 子顺序（按原表行位置填值）
  - 统一列宽、表头字体/填充、数据字体、边框、冻结窗格
  - 标点规范化：半角括号/冒号转全角，连续空格折叠，l/min -> L/min

用法：
  python -I build_reference_format.py --input <zip或目录> --output <输出.xlsx>

可选参数：
  --reference <标准答案.xlsx>   标准答案（第一个工作表，表头含
                               序号/检验工序/检验项目/主顺序/子顺序），不随插件分发
  --project <单项目名称>        xlsx 直接放在根目录时指定项目名
  --keep-temp                   保留 zip 解压临时目录
"""
import argparse
import os
import re
import shutil
import sys
import tempfile
from copy import copy

from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

HERE = os.path.dirname(os.path.abspath(__file__))
if HERE not in sys.path:
    sys.path.insert(0, HERE)

from sip_rules import classify, load_reference, norm, header_key  # noqa: E402
from build_sip_mapping import parse_sheet, discover_projects  # noqa: E402
from extract_gbk_zip import extract as extract_zip  # noqa: E402

try:
    sys.stdout.reconfigure(encoding='utf-8', errors='replace')
except Exception:
    pass

HEADER_FONT = Font(name='微软雅黑', size=10, bold=True, color='FF000000')
DATA_FONT = Font(name='微软雅黑', size=10, color='FF000000')
HEADER_FILL = PatternFill(fill_type='solid', start_color='FFC5D9F1', end_color='FFC5D9F1')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
HEADER_ROW_HEIGHT = 33.75
THIN = Side(style='thin')
WIDTHS = {'A': 11.625, 'B': 11.625, 'C': 11.625, 'D': 13.0, 'E': 13.0,
          'F': 13.0, 'G': 13.0, 'H': 13.0, 'I': 13.0, 'J': 13.0,
          'K': 11.625, 'L': 13.0, 'M': 16.625, 'N': 27.5}


def normalize_text(s):
    if not isinstance(s, str):
        return s
    s = s.replace('(', '（').replace(')', '）').replace(':', '：')
    s = re.sub(r'[ \t]+', ' ', s).strip(' \t')
    return s.replace('l/min', 'L/min')


def find_header_row(ws):
    for r in range(1, min(ws.max_row, 15) + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, ws.max_column + 1)]
        if '序号' in vals and '检验工序' in vals and '检验项目' in vals:
            return r
    raise ValueError('找不到表头行（需要 序号/检验工序/检验项目）')


def transform_source(src_path, row_map):
    """把一个源表转成参考样张格式，返回 Workbook（单 Sheet）。"""
    wb = load_workbook(src_path)
    ws = wb.worksheets[0]
    hr = find_header_row(ws)
    n_title = hr - 1
    max_col0 = ws.max_column
    orig_merges = [(m.min_col, m.min_row, m.max_col, m.max_row)
                   for m in ws.merged_cells.ranges]
    for m in list(ws.merged_cells.ranges):
        ws.unmerge_cells(str(m))

    if n_title:
        ws.delete_rows(1, n_title)

    # 列序号
    col_seq = None
    for c in range(1, ws.max_column + 1):
        if header_key(ws.cell(1, c).value) == '序号':
            col_seq = c
            break
    if col_seq is None:
        raise ValueError('找不到序号列')

    last_seq_row = 1
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, col_seq).value
        if isinstance(v, (int, float)) or (isinstance(v, str) and v.strip().isdigit()):
            last_seq_row = r
    if last_seq_row < ws.max_row:
        ws.delete_rows(last_seq_row + 1, ws.max_row - last_seq_row)
    if max_col0 > 12:
        ws.delete_cols(13, ws.max_column - 12)

    # 标点规范化（在重建合并区域之前）
    for row in ws.iter_rows(min_row=2, max_row=last_seq_row, min_col=1, max_col=12):
        for cell in row:
            cell.value = normalize_text(cell.value)

    # 重建合并区域
    for c1, r1, c2, r2 in orig_merges:
        if c2 > 12:
            continue
        nr1, nr2 = r1 - n_title, r2 - n_title
        if nr2 < 1 or nr1 > last_seq_row:
            continue
        nr1, nr2 = max(1, nr1), min(last_seq_row, nr2)
        if nr2 >= nr1:
            ws.merge_cells(start_row=nr1, start_column=c1, end_row=nr2, end_column=c2)

    # 边框：普通单元格细框；合并区域只画外框；A 列左边 medium；末行下边 medium
    merge_map = {}
    for m in ws.merged_cells.ranges:
        for rr in range(m.min_row, m.max_row + 1):
            for cc in range(m.min_col, m.max_col + 1):
                merge_map[(rr, cc)] = m
    for r in range(2, last_seq_row + 1):
        for c in range(1, 13):
            m = merge_map.get((r, c))
            if m is None:
                left, right, top, bottom = ('medium' if c == 1 else 'thin'), 'thin', 'thin', 'thin'
            else:
                left = 'thin' if c == m.min_col else None
                right = 'thin' if c == m.max_col else None
                top = 'thin' if r == m.min_row else None
                bottom = 'thin' if r == m.max_row else None
                if c == 1 and left:
                    left = 'medium'
            if r == last_seq_row and bottom:
                bottom = 'medium'
            ws.cell(r, c).border = Border(
                left=Side(style=left), right=Side(style=right),
                top=Side(style=top), bottom=Side(style=bottom))

    # M/N
    ws.cell(1, 13, '主顺序')
    ws.cell(1, 14, '子顺序（1级优先进行，12级最后进行）')
    for r in range(2, last_seq_row + 1):
        main, sub = row_map.get(r, ('3. 通电检验', '5级'))
        ws.cell(r, 13, main)
        ws.cell(r, 14, sub)

    # 样式
    for c in range(1, 15):
        cell = ws.cell(1, c)
        cell.font = copy(HEADER_FONT)
        cell.fill = copy(HEADER_FILL)
        cell.alignment = copy(CENTER)
        if c in (13, 14):
            cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
    ws.row_dimensions[1].height = HEADER_ROW_HEIGHT
    for r in range(2, last_seq_row + 1):
        for c in range(1, 15):
            cell = ws.cell(r, c)
            cell.font = copy(DATA_FONT)
            if c in (13, 14):
                cell.alignment = copy(CENTER)
                cell.border = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
        ws.row_dimensions[r].height = None
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A3'
    ws.title = os.path.splitext(os.path.basename(src_path))[0][:31]
    return wb


def copy_sheet(dst_wb, src_ws, title):
    dst = dst_wb.create_sheet(title)
    for row in src_ws.iter_rows():
        for cell in row:
            d = dst.cell(cell.row, cell.column, cell.value)
            if cell.has_style:
                d.font = copy(cell.font)
                d.fill = copy(cell.fill)
                d.border = copy(cell.border)
                d.alignment = copy(cell.alignment)
                d.number_format = cell.number_format
                d.protection = copy(cell.protection)
    for m in src_ws.merged_cells.ranges:
        dst.merge_cells(str(m))
    for key, dim in src_ws.column_dimensions.items():
        dst.column_dimensions[key].width = dim.width
        dst.column_dimensions[key].hidden = dim.hidden
    for r, dim in src_ws.row_dimensions.items():
        dst.row_dimensions[r].height = dim.height
        dst.row_dimensions[r].hidden = dim.hidden
    dst.freeze_panes = src_ws.freeze_panes
    dst.sheet_view.showGridLines = src_ws.sheet_view.showGridLines
    return dst


def safe_name(name, used):
    for ch in '\\/*?:[]':
        name = name.replace(ch, '_')
    name = (name or 'Sheet')[:31]
    base, i = name, 2
    while name in used:
        suffix = f'_{i}'
        name = base[:31 - len(suffix)] + suffix
        i += 1
    used.add(name)
    return name


def verify_reference(out_ws, ref_path):
    ref = load_workbook(ref_path, data_only=True).worksheets[0]
    value_diffs = 0
    for r in range(1, max(ref.max_row, out_ws.max_row) + 1):
        for c in range(1, 15):
            if ref.cell(r, c).value != out_ws.cell(r, c).value:
                value_diffs += 1
    ref_m = set(str(x) for x in ref.merged_cells.ranges)
    out_m = set(str(x) for x in out_ws.merged_cells.ranges)
    return value_diffs, ref_m == out_m


def main():
    ap = argparse.ArgumentParser(description='按参考样张格式生成 SIP 对应表')
    ap.add_argument('--input', required=True, help='zip 或包含 xlsx 的文件夹')
    ap.add_argument('--output', required=True, help='输出 xlsx 路径')
    ap.add_argument('--reference',
                    help='标准答案 xlsx（第一个工作表，表头含 序号/检验工序/检验项目/主顺序/子顺序）')
    ap.add_argument('--project', help='当 xlsx 直接放在根目录时，指定项目名')
    ap.add_argument('--keep-temp', action='store_true', help='保留 zip 解压临时目录')
    args = ap.parse_args()

    tmp = None
    root = args.input
    if os.path.isfile(root) and root.lower().endswith('.zip'):
        tmp = tempfile.mkdtemp(prefix='sip_ref_')
        extract_zip(root, tmp)
        root = tmp
    if not os.path.isdir(root):
        raise SystemExit(f'输入不存在或不是目录/zip: {args.input}')

    ref = load_reference(args.reference) if args.reference else None
    projects = discover_projects(root, args.project)

    out_wb = Workbook()
    out_wb.remove(out_wb.active)
    used = set()
    total = 0
    verify_info = []

    for pname, pdir, files in projects:
        for f in files:
            path = os.path.join(pdir, f)
            rows, _ = parse_sheet(path, pname)
            probe = load_workbook(path, data_only=True)
            n_title = find_header_row(probe.worksheets[0]) - 1
            row_map = {}
            for row in rows:
                (main, sub), _ = classify(
                    row['stage'], row['item'], row['sub'], row['std'], row['method'],
                    ref['pairs'] if ref else None,
                    ref['stages'] if ref else None)
                # 源表行号 -> 转换后行号（减标题行数）
                row_map[row['src'] - n_title] = (main, sub)
            wb = transform_source(path, row_map)
            ws = wb.worksheets[0]
            title = safe_name(ws.title, used)
            dst = copy_sheet(out_wb, ws, title)
            total += max(0, dst.max_row - 1)
            print(f'  {title}: {max(0, dst.max_row - 1)} 行')

            if ref:
                stem = os.path.splitext(f)[0]
                if ref.get('sheet_title') and ref['sheet_title'].lower() in stem.lower():
                    vd, same_merges = verify_reference(dst, args.reference)
                    verify_info.append((stem, vd, same_merges))

    if not out_wb.sheetnames:
        raise SystemExit('没有生成任何 Sheet')
    os.makedirs(os.path.dirname(os.path.abspath(args.output)) or '.', exist_ok=True)
    out_wb.save(args.output)
    print(f'\n已生成: {os.path.abspath(args.output)}')
    print(f'总数据行: {total}')
    for stem, vd, same_merges in verify_info:
        print(f'标准答案比对 {stem}: 值差异 {vd}，合并区域{"一致" if same_merges else "不一致"}')
    if tmp and not args.keep_temp:
        shutil.rmtree(tmp, ignore_errors=True)
    return 0


if __name__ == '__main__':
    sys.exit(main())
